"""Renumber page footers from /08 to /10 to reflect the new 10-page total."""
from pptx import Presentation

PPTX = r"docs\PPT\Agent评测_带封面_v2.pptx"

# (slide_index_0_based, old, new)
RENUMBER = [
    (1, "02 / 08", "02 / 10"),  # P2 Why eval
    (2, "03 / 08", "03 / 10"),  # P3 Compare
    (4, "04 / 08", "05 / 10"),  # P5 Capabilities
    (6, "05 / 08", "07 / 10"),  # P7 Task system
    (7, "06 / 08", "08 / 10"),  # P8 T053
    (8, "07 / 08", "09 / 10"),  # P9 Scenarios
    (9, "08 / 08", "10 / 10"),  # P10 Summary
]

prs = Presentation(PPTX)

for slide_idx, old, new in RENUMBER:
    slide = prs.slides[slide_idx]
    replaced = False
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for para in sh.text_frame.paragraphs:
            full_text = "".join(r.text for r in para.runs)
            if old in full_text:
                new_text = full_text.replace(old, new)
                if para.runs:
                    para.runs[0].text = new_text
                    for r in para.runs[1:]:
                        r.text = ""
                replaced = True
                print(f"  P{slide_idx+1}: {repr(old)} -> {repr(new)}")
                break
        if replaced:
            break
    if not replaced:
        print(f"  WARNING: P{slide_idx+1}: {repr(old)} NOT FOUND")

prs.save(PPTX)
print("saved:", PPTX)
