"""Insert reasons + matrix into the current 8-page Agent评测.pptx (cover already in).

Final order (0-indexed):
  0: Cover           (existing)
  1: P2 Why eval     (existing)
  2: P3 Compare      (existing)
  3: Reasons         (new)
  4: P4 Capabilities (existing)
  5: Matrix          (new)
  6: P5 Task system  (existing)
  7: P6 T053         (existing)
  8: P7 Scenarios    (existing)
  9: P8 Summary      (existing)

Reorder approach: take Reasons first (its position is 8), insert at 3;
then take Matrix (now at 9), insert at 5. Doing it in this order avoids
index shift side effects.
"""
from copy import deepcopy
from pptx import Presentation
from pptx.dml.color import RGBColor

PAGES2_PPTX = r"projects\agentbench_pages2_20260721\exports\agentbench_pages2_20260721_085711.pptx"
ORIG_PPTX = r"docs\PPT\Agent评测.pptx"
OUT = r"docs\PPT\Agent评测_带封面_v2.pptx"

SOURCES = [
    ("Reasons", PAGES2_PPTX, 0),
    ("Matrix", PAGES2_PPTX, 1),
]

prs = Presentation(ORIG_PPTX)
print("original slides:", len(prs.slides._sldIdLst))


def pick_blank_layout(prs):
    for lay in prs.slide_layouts:
        if lay.name.lower() == "blank":
            return lay
    return min(prs.slide_layouts, key=lambda l: len(l.placeholders))


layout = pick_blank_layout(prs)

for label, path, page_idx in SOURCES:
    src = Presentation(path)
    src_slide = src.slides[page_idx]
    new_slide = prs.slides.add_slide(layout)
    for ph in list(new_slide.placeholders):
        el = ph._element
        el.getparent().remove(el)
    bg = new_slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    spTree = new_slide.shapes._spTree
    n = 0
    for shp in list(src_slide.shapes):
        spTree.append(deepcopy(shp._element))
        n += 1
    print(f"  appended {label} ({n} shapes)")

print("after append, total slides:", len(prs.slides._sldIdLst))

# Reorder. After append: [Cover, P2, P3, P4, P5, P6, P7, P8, Reasons, Matrix]
# Indices:                       0     1   2   3   4   5   6   7   8         9
# Step A: move Reasons (idx 8) -> 3
sldIdLst = prs.slides._sldIdLst
reasons_el = list(sldIdLst)[8]
sldIdLst.remove(reasons_el)
sldIdLst.insert(3, reasons_el)
# Now: [Cover, P2, P3, Reasons, P4, P5, P6, P7, P8, Matrix]
# Indices: 0, 1, 2, 3, 4, 5, 6, 7, 8, 9
# Step B: move Matrix (idx 9) -> 5
matrix_el = list(sldIdLst)[9]
sldIdLst.remove(matrix_el)
sldIdLst.insert(5, matrix_el)
# Now: [Cover, P2, P3, Reasons, P4, Matrix, P5, P6, P7, P8]

prs.save(OUT)
print("saved:", OUT)
print("final slides:", len(prs.slides._sldIdLst))
