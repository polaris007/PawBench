# -*- coding: utf-8 -*-
"""在 Agent评测.pptx 第4页(评测任务体系)之后插入一页：以 T053 为例讲解执行/采集/评分/汇总。
保持原 deck 原生形状风格，插入后重排页码为 /08。"""
import re
import copy
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

SRC = r"d:\workplace\github\PawBench\docs\PPT\Agent评测.pptx"
OUT = r"d:\workplace\github\PawBench\docs\PPT\Agent评测.pptx"

# palette
PRIMARY = RGBColor(0x1E, 0x3A, 0x8A)
ACCENT  = RGBColor(0x25, 0x63, 0xEB)
LIGHT   = RGBColor(0xE8, 0xF0, 0xFE)
SOFT    = RGBColor(0xE2, 0xE8, 0xF0)
ALT     = RGBColor(0xF2, 0xF5, 0xFA)
GOOD    = RGBColor(0x10, 0xB9, 0x81)
BAD     = RGBColor(0xEF, 0x44, 0x44)
TEXT    = RGBColor(0x1A, 0x1A, 0x2E)
GRAY    = RGBColor(0x64, 0x74, 0x8B)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)

LATIN = "Segoe UI"
EA = "Microsoft YaHei"


def set_run_font(run, size, bold, color, latin=LATIN, ea=EA):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = latin
    rPr = run._r.get_or_add_rPr()
    # remove existing ea/cs to avoid dup
    for tag in ("a:ea", "a:cs"):
        for el in rPr.findall(qn(tag)):
            rPr.remove(el)
    ea_el = rPr.makeelement(qn("a:ea"), {"typeface": ea})
    cs_el = rPr.makeelement(qn("a:cs"), {"typeface": latin})
    rPr.append(ea_el)
    rPr.append(cs_el)


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, Pt(x), Pt(y), Pt(w), Pt(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(line_w)
    sp.text_frame.paragraphs[0].text = ""
    return sp


def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             line_spacing=1.0, wrap=True):
    """paras: list of paragraphs; each paragraph is list of run dicts {t,size,bold,color}.
    Optional per-paragraph align via first run key 'align'."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    try:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align) if isinstance(para, dict) else align
        if line_spacing:
            p.line_spacing = line_spacing
        runs = para["runs"] if isinstance(para, dict) else para
        if para.get("space_before") if isinstance(para, dict) else None:
            p.space_before = Pt(para["space_before"])
        for rd in runs:
            r = p.add_run()
            r.text = rd["t"]
            set_run_font(r, rd.get("size", 11), rd.get("bold", False), rd.get("color", TEXT),
                         latin=rd.get("latin", LATIN), ea=rd.get("ea", EA))
    return tb


prs = Presentation(SRC)
blank_layout = prs.slides[3].slide_layout
slide = prs.slides.add_slide(blank_layout)

# top accent bar (full width thin) like other pages
add_rect(slide, 0, 0, 960, 3, fill=PRIMARY)

# page number (will be corrected later, but set now)
add_text(slide, 44.6, 16, 90, 16, [[{"t": "06 / 08", "size": 10.5, "bold": False, "color": PRIMARY}]])

# title
add_text(slide, 43.8, 33, 850, 40,
         [[{"t": "任务执行与评分实战 · 以 T053 为例", "size": 30, "bold": True, "color": PRIMARY}]])
# accent underline
add_rect(slide, 45, 72, 60, 2.4, fill=PRIMARY)
# subtitle
add_text(slide, 44.3, 80, 870, 20,
         [[{"t": "用一个真实任务卡，看清 PawBench「执行 → 采集 → 评分 → 汇总」的完整评测链路", "size": 14, "bold": False, "color": GRAY}]])

# ---- Task card ----
add_rect(slide, 45, 104, 870, 58, fill=LIGHT)
add_rect(slide, 45, 104, 4, 58, fill=PRIMARY)
add_text(slide, 60, 111, 850, 20,
         [[{"t": "任务 T053", "size": 12.5, "bold": True, "color": PRIMARY},
           {"t": "   ｜  来源 pinchbench   ｜  类型 内容创作·写作   ｜  复杂度 L1   ｜  环境 closed   ｜  评分 hybrid（自动 0.6 + LLM 0.4）",
            "size": 11.5, "bold": False, "color": TEXT}]])
add_text(slide, 60, 134, 850, 22,
         [[{"t": "指令：", "size": 11, "bold": True, "color": TEXT},
           {"t": "撰写一篇约 500 词《远程办公对软件开发者的好处》博客，并保存为 ", "size": 11, "bold": False, "color": TEXT},
           {"t": "blog_post.md", "size": 11, "bold": True, "color": ACCENT, "latin": "Consolas"},
           {"t": "  （纯文本、无外部依赖）", "size": 11, "bold": False, "color": GRAY}]])

# ---- Four-stage flow ----
CARD_Y = 176
CARD_H = 250
CARD_W = 207
xs = [45, 266, 487, 708]
titles = ["① 沙箱执行", "② 结果采集", "③ 双重评分 · hybrid", "④ 汇总得分"]

# arrows in gaps
for gx in [253, 474, 695]:
    add_text(slide, gx, CARD_Y + 110, 20, 24,
             [[{"t": "▶", "size": 13, "bold": True, "color": ACCENT, "align": PP_ALIGN.CENTER}]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

for i, x in enumerate(xs):
    add_rect(slide, x, CARD_Y, CARD_W, CARD_H, fill=WHITE, line=SOFT, line_w=1.0)
    add_rect(slide, x, CARD_Y, CARD_W, 26, fill=PRIMARY)
    add_text(slide, x, CARD_Y + 4, CARD_W, 20,
             [[{"t": titles[i], "size": 12.5, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

BX = 10  # body inner left pad
BY = CARD_Y + 34
BW = CARD_W - 20

# card 1 body
add_text(slide, xs[0] + BX, BY, BW, 200, [
    [{"t": "Agent 在 Docker 容器中隔离运行，独立执行任务", "size": 10.5, "color": TEXT}],
    {"runs": [{"t": "· 纯文本创作，无需联网", "size": 10.5, "color": TEXT}], "space_before": 8},
    [{"t": "· 无外部依赖 / 无预置文件", "size": 10.5, "color": TEXT}],
    [{"t": "· environment = closed（可离线复现）", "size": 10.5, "color": TEXT}],
    {"runs": [{"t": "· 超时 300s 内产出 ", "size": 10.5, "color": TEXT},
              {"t": "blog_post.md", "size": 10, "bold": True, "color": ACCENT, "latin": "Consolas"}], "space_before": 8},
], line_spacing=1.15)

# card 2 body
add_text(slide, xs[1] + BX, BY, BW, 200, [
    [{"t": "harness 采集运行产物与过程：", "size": 10.5, "color": TEXT}],
    {"runs": [{"t": "· OpenClaw session → 归一化为 transcript 事件流", "size": 10.5, "color": TEXT}], "space_before": 8},
    [{"t": "   (message / toolCall / toolResult)", "size": 9.5, "color": GRAY, "latin": "Consolas"}],
    {"runs": [{"t": "· 新生成文件内容注入 transcript", "size": 10.5, "color": TEXT}], "space_before": 6},
    {"runs": [{"t": "· 产物落盘 ", "size": 10.5, "color": TEXT},
              {"t": "workspace_path", "size": 10, "bold": True, "color": ACCENT, "latin": "Consolas"}], "space_before": 6},
    {"runs": [{"t": "评分只读这两个输入，对框架中立", "size": 9.5, "color": GRAY}], "space_before": 8},
], line_spacing=1.12)

# card 3 body
add_text(slide, xs[2] + BX, BY, BW, 210, [
    [{"t": "自动检查", "size": 10.5, "bold": True, "color": TEXT},
     {"t": "  权重 0.6", "size": 10, "bold": True, "color": PRIMARY}],
    [{"t": "· 文件 blog_post.md 是否存在", "size": 9.5, "color": TEXT}],
    [{"t": "· 词数 450–550 → 1.0（区间打分）", "size": 9.5, "color": TEXT}],
    [{"t": "· 标题 + 段落结构（≥3 段）", "size": 9.5, "color": TEXT}],
    [{"t": "· remote / developer 关键词命中", "size": 9.5, "color": TEXT}],
    {"runs": [{"t": "LLM 评判", "size": 10.5, "bold": True, "color": TEXT},
              {"t": "  权重 0.4", "size": 10, "bold": True, "color": PRIMARY}], "space_before": 8},
    [{"t": "· 质量30 · 结构25 · 写作20", "size": 9.5, "color": TEXT}],
    [{"t": "· 词数15 · 完成度10（Rubric）", "size": 9.5, "color": TEXT}],
], line_spacing=1.1)

# card 4 body
add_text(slide, xs[3] + BX, BY, BW, 210, [
    [{"t": "最终得分", "size": 10.5, "bold": True, "color": TEXT}],
    {"runs": [{"t": "score = 0.6 × 自动分", "size": 10.5, "bold": True, "color": ACCENT}], "space_before": 4},
    [{"t": "            + 0.4 × LLM 分", "size": 10.5, "bold": True, "color": ACCENT}],
    {"runs": [{"t": "混合惩罚机制", "size": 10.5, "bold": True, "color": BAD}], "space_before": 10},
    [{"t": "· 自动分 < 0.75 时", "size": 9.8, "color": TEXT}],
    [{"t": "  → LLM 分被置 0", "size": 9.8, "bold": True, "color": BAD}],
    [{"t": "· 防止跳过任务 / 空文件刷分", "size": 9.5, "color": GRAY}],
    {"runs": [{"t": "(API 真实失败时豁免)", "size": 9, "color": GRAY}], "space_before": 4},
], line_spacing=1.15)

# ---- bottom banner ----
add_rect(slide, 45, 440, 870, 34, fill=PRIMARY)
add_text(slide, 55, 440, 850, 34,
         [[{"t": "一张任务卡 = Prompt + Expected Behavior + 自动检查(代码) + LLM Rubric　｜　评分仅依赖 transcript 与 workspace_path，对 Agent 框架保持中立",
            "size": 12, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---- reorder: move new slide (last) to position index 4 (after old idx3) ----
sldIdLst = prs.slides._sldIdLst
sld_ids = list(sldIdLst)
new_sld = sld_ids[-1]
sldIdLst.remove(new_sld)
sldIdLst.insert(4, new_sld)

# ---- renumber footers: pos -> f"{pos+2:02d} / 08" ----
pat = re.compile(r"^\s*\d{1,2}\s*/\s*\d{1,2}\s*$")
for pos, sld_id in enumerate(list(sldIdLst)):
    rId = sld_id.get(qn("r:id"))
    slide_part = prs.slides._sldIdLst.getparent()  # not used
    s = prs.slides[pos]
    label = f"{pos+2:02d} / 08"
    for sh in s.shapes:
        if sh.has_text_frame and pat.match(sh.text_frame.text or ""):
            tf = sh.text_frame
            # set first run, clear rest
            done = False
            for p in tf.paragraphs:
                for r in p.runs:
                    if not done:
                        r.text = label
                        done = True
                    else:
                        r.text = ""
            break

prs.save(OUT)
print("SAVED", OUT, "slides now", len(prs.slides))
